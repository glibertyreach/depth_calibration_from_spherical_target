"""Check that every text string in the content JSON appears on its slide in the .pptx
(via markitdown), and that every slide's speaker notes equal the JSON notes."""
import json, re, subprocess, sys
from pptx import Presentation

PPTX = "docs/presentations/stage1_capture_overview.pptx"
JSON = "docs/presentations/build/stage1_deck_content.json"
SKIP_KEYS = {"id", "layout", "section", "notes", "icon", "image", "image2", "n"}  # not slide text

def norm(t):
    t = t.replace(" ", " ").replace("\\", "")      # non-breaking space, markdown escapes
    t = re.sub(r"\s+", " ", t)
    return t.strip()

def strings(obj):
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, list):
        for v in obj: yield from strings(v)
    elif isinstance(obj, dict):
        for k, v in obj.items():
            if k not in SKIP_KEYS: yield from strings(v)

md = subprocess.run(["markitdown", PPTX], capture_output=True, text=True, check=True).stdout
parts = re.split(r"<!-- Slide number: (\d+) -->", md)
# markitdown appends "### Notes:" to each slide block; keep the slide body and the notes apart
slide_text, notes_text = {}, {}
for i in range(1, len(parts), 2):
    body, _, notes = parts[i + 1].partition("### Notes:")
    slide_text[int(parts[i])], notes_text[int(parts[i])] = norm(body), norm(notes)
slides = json.load(open(JSON))["slides"]
prs = Presentation(PPTX)
bad = 0
assert len(slides) == len(prs.slides) == len(slide_text), (len(slides), len(prs.slides), len(slide_text))
for i, s in enumerate(slides, 1):
    text = slide_text[i]
    for t in strings({k: v for k, v in s.items()}):
        if norm(t) not in text:
            bad += 1; print(f"slide {i} ({s['id']}): MISSING {t!r}")
    notes = prs.slides[i - 1].notes_slide.notes_text_frame.text if prs.slides[i - 1].has_notes_slide else ""
    if not notes.strip() or norm(notes) != norm(s["notes"]):
        bad += 1; print(f"slide {i} ({s['id']}): NOTES differ or empty")
    # notes text must not be on the slide itself
    if norm(s["notes"]) != notes_text[i]:
        bad += 1; print(f"slide {i}: markitdown notes differ from JSON notes")
    if norm(s["notes"])[:60] in text:
        bad += 1; print(f"slide {i}: notes text found on the slide body")
print(f"checked {len(slides)} slides; problems: {bad}")
sys.exit(1 if bad else 0)
